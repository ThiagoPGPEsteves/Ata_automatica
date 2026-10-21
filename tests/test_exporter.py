"""
Teste Automatizado da Fase 4: Construtor do Documento Word (.docx) com Padrão Executivo Corporativo.
Valida:
1. Geração do arquivo .docx com formatação visual corporativa.
2. Integridade dos elementos: Cabeçalho, Metadados, Caixa de Decisões e Tabela Zebrada de Ações.
3. Pipeline integrado: Markdown do LLM -> MinutesParser -> DocxMeetingMinutesBuilder.
"""

import sys
import os
from pathlib import Path
import docx

# Garante que a raiz do projeto esteja no sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config import OUTPUT_DIR
from src.minutes.parser import MeetingMinutes, ActionItem, MinutesParser
from src.exporter.docx_builder import DocxMeetingMinutesBuilder
from src.exporter.outlook_dispatcher import OutlookDispatcher
import pytest

pytestmark = pytest.mark.unit


def test_docx_generation_from_model():
    """Valida a construção do documento a partir de um objeto MeetingMinutes."""
    print("\n[TESTE 1] Validando Geração de Documento Word (.docx)...")
    minutes = MeetingMinutes(
        title="Reunião de Alinhamento Estratégico e Arquitetura",
        date="14/10/2026",
        participants=["Thiago Esteves", "Coordenação de TI", "Equipe de Engenharia"],
        executive_summary=(
            "A reunião teve como propósito homologar a arquitetura técnica da Ata Automática. "
            "Foram avaliadas as restrições de infraestrutura local, decidindo-se pelo uso de modelos "
            "com quantização int8 e diretiva de descarregamento imediato de memória."
        ),
        topics=[
            {
                "topic": "Arquitetura de Captura de Áudio",
                "details": "Apresentação do módulo WASAPI loopback com gravação contínua em streaming.",
            },
            {
                "topic": "Processamento e Transcrição Local",
                "details": "Homologação do motor faster-whisper com quantização int8 em CPU.",
            },
        ],
        decisions=[
            "Aprovada a implementação em pipeline sequencial desacoplado.",
            "Homologado o limite máximo estrito de 4.0 GB de memória RAM.",
            "Definido o padrão Word (.docx) como formato oficial de distribuição.",
        ],
        action_items=[
            ActionItem(action="Finalizar módulo de exportação Word", owner="Thiago Esteves", deadline="14/10/2026"),
            ActionItem(action="Conectar interface desktop PyQt6", owner="Equipe Frontend", deadline="16/10/2026"),
            ActionItem(action="Executar teste piloto de 30 minutos", owner="Coordenação", deadline="18/10/2026"),
        ],
        next_steps=[
            "Revisão final do código e empacotamento da aplicação.",
            "Apresentação dos resultados da homologação.",
        ],
    )

    output_file = OUTPUT_DIR / "test_ata_executiva.docx"
    builder = DocxMeetingMinutesBuilder(organization_name="Diretoria de Tecnologia / Teams")
    saved_path = builder.build(minutes, output_file)

    assert saved_path.exists(), "O arquivo .docx deveria existir no disco."
    file_size = saved_path.stat().st_size
    print(f"  -> Arquivo gerado com sucesso: {saved_path.name} ({file_size} bytes)")
    assert file_size > 10000, f"Arquivo .docx parece muito pequeno ou vazio: {file_size} bytes"

    print("  -> TESTE 1 APROVADO: Arquivo Word criado com sucesso!")


def test_docx_document_integrity():
    """Valida a estrutura interna do arquivo Word gerado."""
    print("\n[TESTE 2] Validando Integridade Estrutural do Documento Word...")
    test_file = OUTPUT_DIR / "test_ata_executiva.docx"
    assert test_file.exists(), "Arquivo de teste deve existir."

    doc = docx.Document(str(test_file))

    # 1. Valida parágrafos e títulos
    paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
    full_doc_text = " ".join(paragraphs_text)

    assert "ATA DE REUNIÃO EXECUTIVA" in full_doc_text
    assert "Alinhamento Estratégico" in full_doc_text
    assert "1. Resumo Executivo" in full_doc_text
    assert "2. Pautas e Discussões" in full_doc_text
    assert "3. Decisões Tomadas" in full_doc_text
    assert "4. Plano de Ação" in full_doc_text
    assert "5. Próximos Passos" in full_doc_text

    # 2. Valida tabelas (Metadados, Callout de Decisões, Tabela de Ações)
    assert len(doc.tables) >= 3, f"Esperado pelo menos 3 tabelas no documento, encontrado {len(doc.tables)}"

    # Tabela 0: Card de Metadados (2x2)
    meta_table = doc.tables[0]
    assert len(meta_table.rows) == 2
    assert len(meta_table.columns) == 2

    # Tabela 1: Caixa de Decisões (1x1 com decisões dentro)
    decisions_table = doc.tables[1]
    assert len(decisions_table.rows) == 1
    assert "Aprovada a implementação" in decisions_table.cell(0, 0).text

    # Tabela 2: Plano de Ação Zebrado (4 linhas: 1 cabeçalho + 3 dados)
    action_table = doc.tables[2]
    assert len(action_table.rows) == 4, f"Esperado 4 linhas na tabela de ações, encontrado {len(action_table.rows)}"
    assert len(action_table.columns) == 3

    hdr_text = [cell.text.strip() for cell in action_table.rows[0].cells]
    assert "AÇÃO / DELIBERAÇÃO" in hdr_text
    assert "RESPONSÁVEL" in hdr_text
    assert "PRAZO" in hdr_text

    row1_text = [cell.text.strip() for cell in action_table.rows[1].cells]
    assert "Finalizar módulo de exportação Word" in row1_text
    assert "Thiago Esteves" in row1_text

    print(f"  -> Parágrafos validados: {len(paragraphs_text)} parágrafos com títulos e estilos corretos.")
    print(f"  -> Tabelas validadas: {len(doc.tables)} tabelas (Metadados, Decisões e Matriz de Ações).")
    print("  -> TESTE 2 APROVADO: Integridade estrutural do DOCX confirmada!")


def test_markdown_to_docx_pipeline():
    """Valida o fluxo integrado de texto Markdown do LLM para o arquivo .docx."""
    print("\n[TESTE 3] Validando Fluxo Integrado (Markdown -> Parser -> DOCX)...")
    sample_markdown = """# ATA DE REUNIÃO: Homologação do Sistema

## 1. Informações Gerais
- Data: 14/10/2026
- Participantes: Thiago, Gerência

## 2. Resumo Executivo
Sessão de validação da exportação do documento corporativo.

## 3. Pautas Discutidas
### Formatação de Tabelas
Aplicação de paleta azul marinho nas colunas.

## 4. Decisões Tomadas
- Padrão visual homologado sem pendências.

## 5. Plano de Ação
| Ação / O Quê | Responsável | Prazo |
|---|---|---|
| Iniciar desenvolvimento da Fase 5 | Thiago | Imediato |

## 6. Próximos Passos e Pendências
- Conectar widget desktop na próxima fase.
"""

    minutes = MinutesParser.parse(sample_markdown)
    output_pipeline = OUTPUT_DIR / "test_pipeline_output.docx"

    builder = DocxMeetingMinutesBuilder()
    saved = builder.build(minutes, output_pipeline)

    assert saved.exists(), "O arquivo gerado pelo pipeline deveria existir."
    doc = docx.Document(str(saved))
    assert len(doc.tables) >= 3

    print(f"  -> Documento ponta a ponta gerado: {saved.name} ({saved.stat().st_size} bytes)")
    print("  -> TESTE 3 APROVADO: Fluxo integrado concluído com perfeição!")


def test_docx_theme_summary_generation():
    """Valida a exportação de um Resumo Temático sem deliberações ou tabela de ações."""
    print("\n[TESTE 4] Validando Exportação DOCX de Resumo Temático...")
    theme_minutes = MeetingMinutes(
        title="Debate sobre Inteligência Artificial e Governança",
        date="19/10/2026",
        participants=["Thiago", "Sofia", "Carlos"],
        executive_summary="Visão geral e contextualização dos debates sobre o uso ético de LLMs locais.",
        topics=[
            {
                "topic": "Soberania de Dados",
                "details": "Avaliação de privacidade total com modelos locais sem envio de áudio para nuvem externa."
            },
            {
                "topic": "Latência e Desempenho",
                "details": "Métricas de tempo de resposta em CPU mantendo consumo estritamente abaixo de 4GB de RAM."
            }
        ],
        decisions=[],
        action_items=[],
        next_steps=["Acompanhar publicação do novo framework regulatório."],
        document_type="theme_summary",
    )

    out_file = OUTPUT_DIR / "test_resumo_temas.docx"
    builder = DocxMeetingMinutesBuilder()
    saved = builder.build(theme_minutes, out_file)

    assert saved.exists()
    doc = docx.Document(str(saved))

    paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = " ".join(paragraphs_text)

    assert "RESUMO EXECUTIVO DE TEMAS" in full_text
    assert "Debate sobre Inteligência Artificial" in full_text
    assert "1. Visão Geral da Discussão" in full_text
    assert "2. Temas e Tópicos Discutidos" in full_text
    assert "3. Destaques e Observações Relevantes" in full_text
    # Garante que NÃO há cabeçalhos de decisões ou plano de ação vazios
    assert "Plano de Ação" not in full_text
    assert "Decisões Tomadas" not in full_text

    # Deve conter APENAS 1 tabela (o card de metadados). Nenhuma tabela de ações vazia!
    assert len(doc.tables) == 1
    meta_table_text = " ".join([cell.text for row in doc.tables[0].rows for cell in row.cells])
    assert "Resumo Temático de Discussões" in meta_table_text
    print("  -> TESTE 4 APROVADO: Resumo Temático gerado no Word com layout limpo e sem tabelas vazias!")


def test_outlook_dispatcher_subject():
    """Valida a formação do assunto de e-mail para ata e resumo de temas."""
    dispatcher = OutlookDispatcher()

    # 1. Ata Padrão
    minutes = MeetingMinutes(
        title="Alinhamento Semanal de Engenharia",
        date="21/10/2026",
    )
    subj = dispatcher.build_email_subject(minutes)
    assert subj == "Ata de Reunião Executiva: Alinhamento Semanal de Engenharia - 21/10/2026"

    # 2. Resumo de Temas
    theme_minutes = MeetingMinutes(
        title="Debate Geral sobre IA",
        date="21/10/2026",
        document_type="theme_summary",
    )
    theme_subj = dispatcher.build_email_subject(theme_minutes)
    assert theme_subj == "Resumo de Temas: Debate Geral sobre IA - 21/10/2026"


def test_outlook_dispatcher_html_body_and_escaping():
    """Valida a geração de corpo HTML corporativo e escape correto de caracteres especiais."""
    dispatcher = OutlookDispatcher()
    minutes = MeetingMinutes(
        title="Reunião <Estratégica> & Orçamento",
        date="21/10/2026",
        participants=["Thiago <Engenheiro>", "Diretoria & Finanças"],
        executive_summary="Resumo com termos 'críticos' & dados sensíveis.\nSegunda linha.",
        decisions=["Aprovada alocação de 100% dos recursos.", "Implementação em Q1."],
        action_items=[
            ActionItem(action="Contratar consultoria <especial>", owner="Thiago", deadline="30/10/2026")
        ],
        topics=[
            {"title": "Pauta Financeira", "description": "Detalhes de custos e CAPEX."}
        ],
    )

    html_body = dispatcher.build_email_html(minutes, docx_filename="Ata_20261021.docx")

    # 1. Checa escapes
    assert "&lt;Estratégica&gt;" in html_body
    assert "&amp; Orçamento" in html_body
    assert "Thiago &lt;Engenheiro&gt;" in html_body
    assert "&lt;especial&gt;" in html_body

    # 2. Checa estrutura e identidade visual
    assert "#1F4E79" in html_body
    assert "#2E75B6" in html_body
    assert "1. RESUMO EXECUTIVO" in html_body
    assert "2. DELIBERAÇÕES E DECISÕES" in html_body
    assert "3. PLANO DE AÇÃO E RESPONSABILIDADES" in html_body
    assert "Ata_20261021.docx" in html_body
    assert "Documento Oficial Anexado:" in html_body


def test_outlook_dispatcher_create_draft_success_mock(tmp_path):
    """Valida o fluxo bem-sucedido de criação de rascunho com mock do PowerShell COM."""
    from unittest.mock import patch, MagicMock

    fake_docx = tmp_path / "Ata_Teste.docx"
    fake_docx.write_text("fake docx content", encoding="utf-8")

    dispatcher = OutlookDispatcher()
    minutes = MeetingMinutes(
        title="Reunião Piloto",
        date="21/10/2026",
        executive_summary="Resumo de teste.",
    )

    mock_res = MagicMock()
    mock_res.stdout = "SUCCESS:Draft displayed\n"
    mock_res.stderr = ""

    with patch("subprocess.run", return_value=mock_res) as mock_run:
        success, msg = dispatcher.create_outlook_draft(
            minutes=minutes,
            docx_path=fake_docx,
            display=True,
            to_recipients=["equipe@empresa.com", "gerencia@empresa.com"],
        )

        assert success is True
        assert "sucesso" in msg.lower()
        mock_run.assert_called_once()
        cmd_args = mock_run.call_args[0][0]
        assert "powershell" in cmd_args[0]


def test_outlook_dispatcher_create_draft_error_mock(tmp_path):
    """Valida o tratamento de erros do PowerShell COM ao tentar criar o rascunho."""
    from unittest.mock import patch, MagicMock

    fake_docx = tmp_path / "Ata_Teste.docx"
    fake_docx.write_text("fake docx content", encoding="utf-8")

    dispatcher = OutlookDispatcher()
    minutes = MeetingMinutes(title="Reunião Falha", date="21/10/2026")

    mock_res = MagicMock()
    mock_res.stdout = "ERROR:Acesso negado pela política do sistema (GPO).\n"
    mock_res.stderr = ""

    with patch("subprocess.run", return_value=mock_res):
        success, msg = dispatcher.create_outlook_draft(
            minutes=minutes,
            docx_path=fake_docx,
            display=True,
        )

        assert success is False
        assert "Acesso negado" in msg


def test_outlook_dispatcher_missing_docx():
    """Valida que um arquivo .docx inexistente retorna erro imediato sem acionar subprocess."""
    from unittest.mock import patch

    dispatcher = OutlookDispatcher()
    minutes = MeetingMinutes(title="Teste Sem Arquivo")

    with patch("subprocess.run") as mock_run:
        success, msg = dispatcher.create_outlook_draft(
            minutes=minutes,
            docx_path=Path("arquivo_inexistente_123456.docx"),
        )
        assert success is False
        assert "não encontrado" in msg.lower()
        mock_run.assert_not_called()


def test_outlook_dispatcher_mailto_fallback(tmp_path):
    """Valida o funcionamento do fallback via mailto e abertura de pasta."""
    from unittest.mock import patch

    fake_docx = tmp_path / "Ata_Fallback.docx"
    fake_docx.write_text("dummy", encoding="utf-8")

    dispatcher = OutlookDispatcher()
    minutes = MeetingMinutes(title="Reunião Fallback", date="21/10/2026", executive_summary="Resumo")

    with patch("webbrowser.open") as mock_web, patch("os.startfile") as mock_startfile:
        res = dispatcher.open_mailto_fallback(minutes, fake_docx)
        assert res is True
        mock_web.assert_called_once()
        assert "mailto:" in mock_web.call_args[0][0]
        mock_startfile.assert_called_once_with(str(fake_docx.parent.resolve()))


if __name__ == "__main__":
    test_docx_generation_from_model()
    test_docx_document_integrity()
    test_markdown_to_docx_pipeline()
    test_docx_theme_summary_generation()
    test_outlook_dispatcher_subject()
    test_outlook_dispatcher_html_body_and_escaping()
    test_outlook_dispatcher_create_draft_success_mock(Path("."))
    test_outlook_dispatcher_create_draft_error_mock(Path("."))
    test_outlook_dispatcher_missing_docx()
    test_outlook_dispatcher_mailto_fallback(Path("."))
    print("\n==========================================")
    print("TODOS OS TESTES DA FASE 4 FORAM CONCLUÍDOS!")
    print("==========================================")
