"""
Construtor de Documentos Microsoft Word (.docx) para Atas de Reunião
Gera documentos corporativos elegantes com cabeçalho executivo, caixas de destaque
para deliberações e tabelas zebradas para o plano de ação.
"""

from pathlib import Path
from typing import Union, Optional
from datetime import datetime
import logging

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from src.minutes.parser import MeetingMinutes
from src.exporter.styles import (
    CorporatePalette,
    DocumentStyles,
    set_cell_background,
    set_cell_margins,
    set_callout_border,
    set_table_borders,
)

logger = logging.getLogger(__name__)


class DocxMeetingMinutesBuilder:
    """
    Constrói um arquivo .docx formatado com padrão visual executivo
    a partir de um objeto MeetingMinutes estruturado.
    """

    def __init__(self, organization_name: str = "Corporativo / Microsoft Teams"):
        self.organization_name = organization_name

    def build(
        self,
        minutes: MeetingMinutes,
        output_filepath: Union[str, Path],
    ) -> Path:
        """Gera o documento Word completo e salva no caminho indicado."""
        out_path = Path(output_filepath)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # 1. Configura margens do documento (1 polegada = 2.54 cm)
        sections = doc.sections
        for section in sections:
            section.top_margin = DocumentStyles.PAGE_MARGIN
            section.bottom_margin = DocumentStyles.PAGE_MARGIN
            section.left_margin = DocumentStyles.PAGE_MARGIN
            section.right_margin = DocumentStyles.PAGE_MARGIN

            # Rodapé corporativo
            footer = section.footer
            footer_para = footer.paragraphs[0]
            footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            f_run = footer_para.add_run(
                f"Documento Confidencial | Ata de Reunião Teams | Gerado em {datetime.now().strftime('%d/%m/%Y')}"
            )
            f_run.font.name = DocumentStyles.FONT_FAMILY
            f_run.font.size = DocumentStyles.FOOTER_SIZE
            f_run.font.color.rgb = CorporatePalette.TEXT_MUTED

        # 2. Cabeçalho Executivo
        self._add_title_block(doc, minutes)

        # 3. Metadados da Reunião (Data, Participantes)
        self._add_metadata_card(doc, minutes)

        if minutes.is_theme_summary:
            # Modo Resumo Temático
            sec_idx = 1
            # 4. Visão Geral
            if minutes.executive_summary:
                self._add_section_header(doc, f"{sec_idx}. Visão Geral da Discussão")
                sec_idx += 1
                para = doc.add_paragraph()
                para.paragraph_format.line_spacing = 1.15
                para.paragraph_format.space_after = Pt(10)
                run = para.add_run(minutes.executive_summary)
                run.font.name = DocumentStyles.FONT_FAMILY
                run.font.size = DocumentStyles.BODY_SIZE
                run.font.color.rgb = CorporatePalette.TEXT_DARK

            # 5. Temas e Tópicos Discutidos
            if minutes.topics:
                self._add_section_header(doc, f"{sec_idx}. Temas e Tópicos Discutidos")
                sec_idx += 1
                for t in minutes.topics:
                    p_topic = doc.add_paragraph()
                    p_topic.paragraph_format.space_before = Pt(6)
                    p_topic.paragraph_format.space_after = Pt(3)
                    r_topic = p_topic.add_run(f"▸ {t.get('topic', 'Tema')}")
                    r_topic.bold = True
                    r_topic.font.name = DocumentStyles.FONT_FAMILY
                    r_topic.font.size = DocumentStyles.SUBSECTION_SIZE
                    r_topic.font.color.rgb = CorporatePalette.STEEL_BLUE

                    if t.get("details"):
                        p_det = doc.add_paragraph()
                        p_det.paragraph_format.left_indent = Inches(0.2)
                        p_det.paragraph_format.line_spacing = 1.15
                        p_det.paragraph_format.space_after = Pt(6)
                        r_det = p_det.add_run(t["details"])
                        r_det.font.name = DocumentStyles.FONT_FAMILY
                        r_det.font.size = DocumentStyles.BODY_SIZE
                        r_det.font.color.rgb = CorporatePalette.TEXT_DARK

            # 6. Destaques e Observações
            if minutes.next_steps:
                self._add_section_header(doc, f"{sec_idx}. Destaques e Observações Relevantes")
                sec_idx += 1
                for step in minutes.next_steps:
                    p_step = doc.add_paragraph()
                    p_step.paragraph_format.left_indent = Inches(0.15)
                    p_step.paragraph_format.space_after = Pt(4)
                    r_bullet = p_step.add_run("▪ ")
                    r_bullet.bold = True
                    r_bullet.font.color.rgb = CorporatePalette.STEEL_BLUE
                    r_text = p_step.add_run(step)
                    r_text.font.name = DocumentStyles.FONT_FAMILY
                    r_text.font.size = DocumentStyles.BODY_SIZE
                    r_text.font.color.rgb = CorporatePalette.TEXT_DARK

            # Deliberações e Ações (apenas se existirem explicitamente no resumo)
            if minutes.decisions:
                self._add_section_header(doc, f"{sec_idx}. Deliberações Mencionadas")
                sec_idx += 1
                self._add_decisions_callout(doc, minutes.decisions)

            if minutes.action_items:
                self._add_section_header(doc, f"{sec_idx}. Ações Mencionadas")
                sec_idx += 1
                self._add_action_items_table(doc, minutes.action_items)

        else:
            # Modo Ata Executiva Padrão
            # 4. Resumo Executivo
            if minutes.executive_summary:
                self._add_section_header(doc, "1. Resumo Executivo")
                para = doc.add_paragraph()
                para.paragraph_format.line_spacing = 1.15
                para.paragraph_format.space_after = Pt(10)
                run = para.add_run(minutes.executive_summary)
                run.font.name = DocumentStyles.FONT_FAMILY
                run.font.size = DocumentStyles.BODY_SIZE
                run.font.color.rgb = CorporatePalette.TEXT_DARK

            # 5. Pautas e Discussões
            if minutes.topics:
                self._add_section_header(doc, "2. Pautas e Discussões")
                for t in minutes.topics:
                    # Subtítulo da pauta
                    p_topic = doc.add_paragraph()
                    p_topic.paragraph_format.space_before = Pt(6)
                    p_topic.paragraph_format.space_after = Pt(3)
                    r_topic = p_topic.add_run(f"▸ {t.get('topic', 'Pauta')}")
                    r_topic.bold = True
                    r_topic.font.name = DocumentStyles.FONT_FAMILY
                    r_topic.font.size = DocumentStyles.SUBSECTION_SIZE
                    r_topic.font.color.rgb = CorporatePalette.STEEL_BLUE

                    # Detalhes da pauta
                    if t.get("details"):
                        p_det = doc.add_paragraph()
                        p_det.paragraph_format.left_indent = Inches(0.2)
                        p_det.paragraph_format.line_spacing = 1.15
                        p_det.paragraph_format.space_after = Pt(6)
                        r_det = p_det.add_run(t["details"])
                        r_det.font.name = DocumentStyles.FONT_FAMILY
                        r_det.font.size = DocumentStyles.BODY_SIZE
                        r_det.font.color.rgb = CorporatePalette.TEXT_DARK

            # 6. Decisões Tomadas (Caixa de Destaque / Callout Box)
            if minutes.decisions:
                self._add_section_header(doc, "3. Decisões Tomadas e Deliberações")
                self._add_decisions_callout(doc, minutes.decisions)

            # 7. Plano de Ação (Tabela Zebrada Corporativa)
            if minutes.action_items:
                self._add_section_header(doc, "4. Plano de Ação (Matriz de Responsabilidades)")
                self._add_action_items_table(doc, minutes.action_items)

            # 8. Próximos Passos e Pendências
            if minutes.next_steps:
                self._add_section_header(doc, "5. Próximos Passos e Encaminhamentos")
                for step in minutes.next_steps:
                    p_step = doc.add_paragraph()
                    p_step.paragraph_format.left_indent = Inches(0.15)
                    p_step.paragraph_format.space_after = Pt(4)
                    r_bullet = p_step.add_run("▪ ")
                    r_bullet.bold = True
                    r_bullet.font.color.rgb = CorporatePalette.STEEL_BLUE
                    r_text = p_step.add_run(step)
                    r_text.font.name = DocumentStyles.FONT_FAMILY
                    r_text.font.size = DocumentStyles.BODY_SIZE
                    r_text.font.color.rgb = CorporatePalette.TEXT_DARK

        doc.save(str(out_path))
        logger.info("Documento Word (.docx) gerado com sucesso: %s", out_path)
        return out_path

    def _add_title_block(self, doc, minutes: MeetingMinutes):
        """Adiciona o título principal estilizado com faixa e subtítulo."""
        p_org = doc.add_paragraph()
        p_org.paragraph_format.space_after = Pt(2)
        r_org = p_org.add_run(self.organization_name.upper())
        r_org.font.name = DocumentStyles.FONT_FAMILY
        r_org.font.size = Pt(8.5)
        r_org.bold = True
        r_org.font.color.rgb = CorporatePalette.STEEL_BLUE

        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_after = Pt(4)
        doc_header = "RESUMO EXECUTIVO DE TEMAS" if minutes.is_theme_summary else "ATA DE REUNIÃO EXECUTIVA"
        r_title = p_title.add_run(doc_header)
        r_title.font.name = DocumentStyles.FONT_FAMILY
        r_title.font.size = DocumentStyles.TITLE_SIZE
        r_title.bold = True
        r_title.font.color.rgb = CorporatePalette.NAVY

        default_names = ("Ata de Reunião", "Resumo de Temas")
        if minutes.title and minutes.title not in default_names:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_after = Pt(12)
            r_sub = p_sub.add_run(minutes.title)
            r_sub.font.name = DocumentStyles.FONT_FAMILY
            r_sub.font.size = Pt(13)
            r_sub.font.color.rgb = CorporatePalette.STEEL_BLUE

    def _add_metadata_card(self, doc, minutes: MeetingMinutes):
        """Cria um card discreto de metadados (Data, Participantes)."""
        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        date_val = minutes.date or datetime.now().strftime("%d/%m/%Y")
        participants_val = ", ".join(minutes.participants) if minutes.participants else "Identificados por canal (Você / Equipe)"
        status_label = "TIPO DE DOCUMENTO" if minutes.is_theme_summary else "SITUAÇÃO DA ATA"
        status_val = "Resumo Temático de Discussões" if minutes.is_theme_summary else "Aprovada para Distribuição"

        meta_items = [
            ("DATA DA REUNIÃO", date_val),
            ("PLATAFORMA", "Microsoft Teams (Gravação Dual-Channel)"),
            ("PARTICIPANTES", participants_val),
            (status_label, status_val),
        ]

        idx = 0
        for r in range(2):
            for c in range(2):
                cell = table.cell(r, c)
                set_cell_background(cell, CorporatePalette.HEX_LIGHT_BG)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                label, val = meta_items[idx]
                idx += 1

                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                r_lbl = p.add_run(f"{label}: ")
                r_lbl.bold = True
                r_lbl.font.size = Pt(8.5)
                r_lbl.font.color.rgb = CorporatePalette.NAVY

                r_val = p.add_run(val)
                r_val.font.size = Pt(9.0)
                r_val.font.color.rgb = CorporatePalette.TEXT_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def _add_section_header(self, doc, title: str):
        """Adiciona cabeçalho de seção numerada em azul marinho."""
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.bold = True
        run.font.name = DocumentStyles.FONT_FAMILY
        run.font.size = DocumentStyles.SECTION_SIZE
        run.font.color.rgb = CorporatePalette.NAVY

    def _add_decisions_callout(self, doc, decisions: list):
        """Renderiza decisões em caixa de destaque com borda lateral espessa."""
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)

        set_cell_background(cell, CorporatePalette.HEX_DECISION_BG)
        set_callout_border(cell, CorporatePalette.HEX_DECISION_BORDER)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        for i, dec in enumerate(decisions):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r_icon = p.add_run("✔ ")
            r_icon.bold = True
            r_icon.font.color.rgb = CorporatePalette.STEEL_BLUE
            r_text = p.add_run(dec)
            r_text.bold = True
            r_text.font.name = DocumentStyles.FONT_FAMILY
            r_text.font.size = DocumentStyles.BODY_SIZE
            r_text.font.color.rgb = CorporatePalette.TEXT_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def _add_action_items_table(self, doc, action_items: list):
        """Cria tabela formatada e zebrada para a matriz de ações."""
        table = doc.add_table(rows=len(action_items) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        col_widths = [Inches(3.3), Inches(1.7), Inches(1.4)]

        # Cabeçalho da Tabela
        headers = ["AÇÃO / DELIBERAÇÃO", "RESPONSÁVEL", "PRAZO"]
        hdr_row = table.rows[0]
        for c_idx, title in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, CorporatePalette.HEX_TABLE_HEADER)
            set_cell_margins(cell, top=120, bottom=120, left=120, right=120)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(title)
            run.bold = True
            run.font.name = DocumentStyles.FONT_FAMILY
            run.font.size = DocumentStyles.TABLE_SIZE
            run.font.color.rgb = CorporatePalette.WHITE

        # Linhas de Dados Zebradas
        for r_idx, item in enumerate(action_items, start=1):
            row = table.rows[r_idx]
            bg_color = CorporatePalette.HEX_ZEBRA_ODD if r_idx % 2 == 1 else CorporatePalette.HEX_ZEBRA_EVEN

            act_text = item.action
            ev = getattr(item, "evidence_ts", "")
            if ev and f"[{ev}]" not in act_text and f"({ev})" not in act_text and f"ref: {ev}" not in act_text:
                act_text = f"{act_text} (ref: {ev})"

            values = [act_text, item.owner, item.deadline]
            for c_idx, val in enumerate(values):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(val)
                run.font.name = DocumentStyles.FONT_FAMILY
                run.font.size = DocumentStyles.TABLE_SIZE
                run.font.color.rgb = CorporatePalette.TEXT_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(8)
